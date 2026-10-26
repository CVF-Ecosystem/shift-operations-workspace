"""Authenticated export across both P2-R ledgers; no ledger mutation."""

import hashlib
import sys
from datetime import UTC, datetime, timedelta
from io import BytesIO
from pathlib import Path
from uuid import uuid4

import pytest
from _auth_test_helpers import auth_headers
from cvf_runtime.identity import Principal
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from operations_domain.models import Shift
from operations_domain.report_digest import canonical_bytes
from operations_ledger.sql_ledger import SqlLedger, make_engine
from operations_ledger.tables import metadata
from workspace_api.application.report_service import ReportService
from workspace_api.dependencies import get_ledger
from workspace_api.domain import models as domain_models
from workspace_api.infrastructure.repository import InMemoryLedger
from workspace_api.main import app

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "scripts"))
from invariant_family_contract import generate_mutations, load_json_no_dup, matches_shape_exactly

MATRIX = load_json_no_dup(ROOT / "docs/cvf/invariants/p5a-export-outcomes.json")


def _assert_refusal(response, outcome: str, before: str, ledger, report_id) -> None:
    shape = next(s for item in MATRIX["outcomes"] if item["outcomeId"] == outcome
                 for s in item["shapes"])
    assert "content-disposition" not in response.headers
    assert response.headers["content-type"].startswith("application/json")
    receipt = {
        "outcome": outcome, "http_status": response.status_code,
        "attachment": False, "binary_bytes": 0,
        "before_state_digest": before,
        "after_state_digest": hashlib.sha256(canonical_bytes(
            ledger.get_report(report_id).model_dump(mode="json"),
        )).hexdigest(),
    }
    if "error_code" in shape["requiredFields"]:
        receipt["error_code"] = response.json()["detail"]["code"]
    assert matches_shape_exactly(shape, receipt, MATRIX)
    mutations = generate_mutations(MATRIX, shape["shapeId"], receipt)
    assert mutations
    assert all(not matches_shape_exactly(shape, m.payload, MATRIX) for m in mutations)


@pytest.mark.parametrize("backend", ["memory", "sqlite"])
def test_api_download_assignment_refusals_and_nonmutation(tmp_path, backend, monkeypatch):
    if backend == "memory":
        ledger = InMemoryLedger()
    else:
        path = tmp_path / "reports.sqlite3"
        engine = make_engine(f"sqlite:///{path}")
        metadata.create_all(engine)
        ledger = SqlLedger(str(path), models=domain_models, engine=engine)
    now = datetime.now(UTC)
    shift = Shift(name="Day", starts_at=now, ends_at=now + timedelta(hours=8))
    ledger.create_shift(shift)
    ledger.add_user(domain_models.User(user_id="op1", username="op1", password_hash="x", role="operator"))
    ledger.add_assignment(domain_models.ShiftAssignment(
        shift_id=shift.shift_id, user_id="op1", assigned_by="op1"
    ))
    ledger.close_shift(shift.shift_id)
    report = ReportService(ledger).generate(shift.shift_id, Principal(user_id="op1", role="operator"))
    before = ledger.get_report(report.report_id).model_dump(mode="json")
    before_digest = hashlib.sha256(canonical_bytes(before)).hexdigest()
    app.dependency_overrides[get_ledger] = lambda: ledger
    try:
        client = TestClient(app)
        path = f"/reports/{report.report_id}/export"
        query = {"format": "pdf", "view": "shift-summary"}
        unauthorized = client.get(path, params=query)
        assert unauthorized.status_code == 401
        _assert_refusal(unauthorized, "AUTH_REFUSED", before_digest, ledger, report.report_id)
        inaccessible = client.get(path, params=query, headers=auth_headers("other", "operator"))
        assert inaccessible.status_code == 404
        _assert_refusal(inaccessible, "NOT_FOUND", before_digest, ledger, report.report_id)
        invalid = client.get(path, params={"format": "bad", "view": "bad"}, headers=auth_headers("op1", "operator"))
        assert invalid.status_code == 422
        assert invalid.json()["detail"]["code"] == "EXPORT_INVALID_PARAMETER"
        _assert_refusal(invalid, "INVALID_REQUEST", before_digest, ledger, report.report_id)
        good = client.get(path, params=query, headers=auth_headers("op1", "operator"))
        assert good.status_code == 200 and good.content.startswith(b"%PDF")
        assert good.headers["cache-control"] == "no-store"
        assert good.headers["x-content-type-options"] == "nosniff"
        assert "attachment;" in good.headers["content-disposition"]
        xlsx = client.get(path, params={"format": "xlsx", "view": "shift-summary"},
                          headers=auth_headers("op1", "operator"))
        assert xlsx.status_code == 200
        assert xlsx.headers["content-type"].startswith(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        assert xlsx.headers["cache-control"] == "no-store"
        assert "attachment;" in xlsx.headers["content-disposition"]
        wb = load_workbook(BytesIO(xlsx.content), data_only=False)
        stored = ledger.get_report(report.report_id)
        iso = stored.generated_from_cutoff.isoformat()
        assert dict(list(wb["Metadata"].values)[1:])["generated_from_cutoff"] == iso
        assert list(wb["Date values"].values)[1][2] == iso
        note = list(wb["Date values"].values)[1][4]
        if backend == "sqlite":
            assert stored.generated_from_cutoff.tzinfo is None
            assert "Timezone unspecified" in note and "no conversion" in note
        else:
            assert "converted to UTC" in note
        assert wb["Metadata"]["B5"].data_type == "n"  # persisted version
        from workspace_api.api.reports import router as export_router
        from workspace_api.application import report_export

        assert report_export._render_slot.acquire(blocking=False)
        try:
            busy = client.get(path, params=query, headers=auth_headers("op1", "operator"))
            assert busy.status_code == 503
            assert busy.json()["detail"]["code"] == "EXPORT_BUSY"
            assert "content-disposition" not in busy.headers
            _assert_refusal(busy, "RENDER_BUSY", before_digest, ledger, report.report_id)
        finally:
            report_export._render_slot.release()

        def fail_unexpected(*args):
            raise RuntimeError("raw operational contents should never be disclosed")

        monkeypatch.setattr(export_router, "export_report", fail_unexpected)
        failed = client.get(path, params=query, headers=auth_headers("op1", "operator"))
        assert failed.status_code == 500
        assert failed.json()["detail"]["code"] == "EXPORT_FAILED"
        assert "raw operational contents" not in failed.text
        assert "content-disposition" not in failed.headers
        _assert_refusal(failed, "RENDER_FAILED", before_digest, ledger, report.report_id)
        assert ledger.get_report(report.report_id).model_dump(mode="json") == before
        missing = client.get(f"/reports/{uuid4()}/export", params=query,
                             headers=auth_headers("op1", "operator"))
        assert missing.status_code == 404
        _assert_refusal(missing, "NOT_FOUND", before_digest, ledger, report.report_id)
    finally:
        app.dependency_overrides.pop(get_ledger, None)
