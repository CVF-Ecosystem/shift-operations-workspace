"""Independent six-section input oracle and inspectable PDF/XLSX artifacts."""

import hashlib
import json
import re
import sys
from collections import Counter
from datetime import UTC, datetime
from io import BytesIO
from pathlib import Path
from uuid import uuid4

import fitz
import pytest
from openpyxl import load_workbook
from operations_domain.models import (
    Report,
    ReportContent,
    ReportSection,
    ReportSourceRef,
    ReportStatus,
)
from operations_domain.report_digest import aggregate_digest, canonical_bytes
from reporting_engine import render_bounded

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "scripts"))
from invariant_family_contract import generate_mutations, load_json_no_dup, matches_shape_exactly

PROOF_DIR = ROOT / ".venv" / "p5a-worker-proof"
MATRIX = load_json_no_dup(ROOT / "docs/cvf/invariants/p5a-export-outcomes.json")
SECTION_TYPES = (
    "operational_events", "corrections", "tasks", "customer_requests", "incidents", "handovers",
)
STATUS_LABELS = {
    "DRAFT": "Draft", "IN_REVIEW": "In review", "APPROVED": "Approved", "FROZEN": "Frozen",
}
TIME = "2026-01-01T00:00:00Z"


def _membership_digest(members) -> str:
    return hashlib.sha256(canonical_bytes(sorted(members))).hexdigest()


def _parsed_detail_refs(pairs) -> list[tuple[str, ...]]:
    refs = []
    for key, value in pairs:
        if key == "Evidence":
            item = json.loads(value)
            refs.append(("evidence", item["record_type"], item["record_id"],
                         item["item_id"] or "", item["evidence_id"], item["source_type"],
                         item["source_id"], item["sha256"] or ""))
        elif key == "Handover source":
            item = json.loads(value)
            refs.append(("handover", item["handover_id"], item["item_id"],
                         item["source_record_type"], item["source_record_id"],
                         item["source_digest"]))
    return refs


def _pdf_json_rows(flat: str, label: str) -> list[tuple[str, str]]:
    decoder = json.JSONDecoder()
    found = []
    start = 0
    needle = label + ": "
    while (position := flat.find(needle, start)) >= 0:
        item, length = decoder.raw_decode(flat[position + len(needle):])
        found.append((label, json.dumps(item)))
        start = position + len(needle) + length
    return found


def _metadata_digest(values: dict[str, str]) -> str:
    return hashlib.sha256(canonical_bytes(values)).hexdigest()


def _evidence():
    return [{"evidence_id": str(uuid4()), "source_type": "manual", "source_id": "E-1", "sha256": None}]


def full_report() -> Report:
    shift = uuid4()
    event, correction, task, request, incident, handover = [str(uuid4()) for _ in range(6)]
    evidence = _evidence()
    base = {
        "operational_events": {"record_type": "OperationalEvent", "record_id": event,
            "shift_id": str(shift), "event_type": "NOTE", "title": "Báo cáo", "description": "Đã xác nhận",
            "risk_class": "R1", "state": "CONFIRMED", "starts_at": TIME, "ends_at": None,
            "owner_id": "op1", "version": 1, "evidence": evidence},
        "corrections": {"record_type": "Correction", "record_id": correction,
            "target_record_type": "OperationalEvent", "target_record_id": event,
            "reason": "Đính chính", "requested_by": "op1", "previous_version": 1,
            "new_version": 2, "created_at": TIME},
        "tasks": {"record_type": "Task", "record_id": task, "shift_id": str(shift),
            "title": "Báo cáo", "description": "=HYPERLINK(\"https://example.test\") " + "x" * 800,
            "status": "OPEN", "owner_id": "op1", "due_at": None, "risk_class": "R1",
            "state": "CONFIRMED", "version": 1, "created_at": TIME, "evidence": evidence},
        "customer_requests": {"record_type": "CustomerRequest", "record_id": request,
            "customer_id": "khách", "shift_id": str(shift), "summary": "Yêu cầu",
            "details": "literal + @ - tab\tvalue", "status": "NEW", "source_message_id": None,
            "received_at": TIME, "promised_at": None, "owner_id": "op1"},
        "incidents": {"record_type": "Incident", "record_id": incident,
            "shift_id": str(shift), "risk_class": "R2", "summary": "Sự cố",
            "description": "<script>literal</script>", "status": "REPORTED", "owner_id": "op1",
            "version": 1, "created_at": TIME, "evidence": evidence},
        "handovers": {"record_type": "Handover", "record_id": handover,
            "from_shift_id": str(shift), "to_shift_id": str(uuid4()), "status": "DRAFT",
            "created_by": "op1", "reviewed_by": None, "reviewed_at": None,
            "received_by": None, "acknowledged_at": None, "version": 1,
            "created_at": TIME, "items": [{"item_id": str(uuid4()),
                "source_record_type": "Task", "source_record_id": task,
                "source_digest": "a" * 64, "summary": "Chuyển giao",
                "owner_id": "op1", "due_at": None, "risk_class": "R1", "evidence": evidence}]},
    }
    sections = [ReportSection(section_type=name, records=[base[name]]) for name in SECTION_TYPES]
    refs = [ReportSourceRef(record_type=record["record_type"], record_id=record["record_id"],
              source_version=record.get("version"),
              source_digest=hashlib.sha256(canonical_bytes(record)).hexdigest())
            for record in base.values()]
    content = ReportContent(sections=sections, source_manifest=refs, snapshot_digest="0" * 64)
    content.snapshot_digest = aggregate_digest(shift, "END_SHIFT", content)
    return Report(shift_id=shift, content=content,
                  generated_from_cutoff=datetime(2026, 1, 1, tzinfo=UTC))


@pytest.mark.parametrize("status", list(STATUS_LABELS))
@pytest.mark.parametrize("current", [True, False])
def test_all_status_history_full_artifacts(status, current):
    report = full_report()
    report.status = ReportStatus(status)
    report.is_current = current
    expected_records = [(r["record_type"], r["record_id"])
                        for s in report.content.sections for r in s.records]
    expected_manifest = {(m.record_type, str(m.record_id), str(m.source_version), m.source_digest)
                         for m in report.content.source_manifest}
    expected_refs = [("manifest", *ref) for ref in expected_manifest]
    for section in report.content.sections:
        for record in section.records:
            for evidence in record.get("evidence", []):
                expected_refs.append(("evidence", record["record_type"], record["record_id"],
                                      "", evidence["evidence_id"], evidence["source_type"],
                                      evidence["source_id"], evidence["sha256"] or ""))
            for item in record.get("items", []):
                expected_refs.append(("handover", record["record_id"], item["item_id"],
                                      item["source_record_type"], item["source_record_id"],
                                      item["source_digest"]))
                for evidence in item.get("evidence", []):
                    expected_refs.append(("evidence", record["record_type"], record["record_id"],
                                          item["item_id"], evidence["evidence_id"],
                                          evidence["source_type"], evidence["source_id"],
                                          evidence["sha256"] or ""))
    assert sum(ref[0] == "evidence" for ref in expected_refs) == 4
    assert sum(ref[0] == "handover" for ref in expected_refs) == 1
    expected_metadata = {
        "report_id": str(report.report_id), "shift_id": str(report.shift_id),
        "report_type": "END_SHIFT", "version": str(report.version),
        "status": STATUS_LABELS[status], "is_current": "Yes" if current else "No",
        "generated_from_cutoff": report.generated_from_cutoff.isoformat(),
        "snapshot_digest": report.content.snapshot_digest,
        "view": "shift-summary", "template_version": "p5a-1",
    }
    if not current:
        expected_metadata["Historical"] = "Yes"
    before = hashlib.sha256(canonical_bytes(report.model_dump(mode="json"))).hexdigest()
    PROOF_DIR.mkdir(parents=True, exist_ok=True)
    name = f"full-{status.lower()}-{'current' if current else 'historical'}"
    pdf = render_bounded(report, "shift-summary", "pdf")
    xlsx = render_bounded(report, "shift-summary", "xlsx")
    (PROOF_DIR / f"{name}.pdf").write_bytes(pdf)
    (PROOF_DIR / f"{name}.xlsx").write_bytes(xlsx)
    document = fitz.open(stream=pdf, filetype="pdf")
    text = "\n".join(p.get_text() for p in document)
    for page in list(document)[1:]:
        assert any(line.startswith("Continued section:") for line in page.get_text().splitlines()[:5])
    assert any("Continued record:" in page.get_text() for page in list(document)[1:])
    flat = "".join(line for line in text.splitlines() if not line.startswith((
        "End-shift report", "Page ", "Continued section:", "Continued record:",
    )))
    parsed_records = re.findall(r"Record: (\w+)/([0-9a-f-]{36})", text)
    assert Counter(parsed_records) == Counter(expected_records)
    parsed_manifest = re.findall(
        r"Source: (\w+)/([0-9a-f-]{36})/([^/]+)/([0-9a-f]{64})", flat,
    )
    assert Counter(parsed_manifest) == Counter(expected_manifest)
    parsed_refs = [("manifest", *ref) for ref in parsed_manifest]
    parsed_refs += _parsed_detail_refs(_pdf_json_rows(flat, "Evidence") +
                                        _pdf_json_rows(flat, "Handover source"))
    assert Counter(parsed_refs) == Counter(expected_refs)
    assert f"status: {STATUS_LABELS[status]}" in text
    assert f"is_current: {'Yes' if current else 'No'}" in text
    assert ("Historical: Yes" in text) is (not current)
    assert f"report_id: {report.report_id}" in text
    assert f"shift_id: {report.shift_id}" in text
    assert "report_type: END_SHIFT" in text
    assert f"snapshot_digest: {report.content.snapshot_digest}" in text
    assert "template_version: p5a-1" in text
    assert "Báo cáo" in text and "Đã xác nhận" in text
    assert "Báo cáo" not in text  # PDF presentation normalizes NFD to NFC.
    assert "<script>literal</script>" in text
    assert "tab\\tvalue" in text
    pdf_meta = {key: re.search(rf"^{key}: (.+)$", text, re.MULTILINE).group(1)
                for key in expected_metadata}
    assert pdf_meta == expected_metadata
    wb = load_workbook(BytesIO(xlsx), data_only=False)
    meta = dict(list(wb["Metadata"].values)[1:])
    assert meta["status"] == STATUS_LABELS[status]
    assert meta["is_current"] == ("Yes" if current else "No")
    assert (meta.get("Historical") == "Yes") is (not current)
    cells = list(wb["Content"].values)
    assert Counter((value.split("/", 1)[0], value.split("/", 1)[1])
                   for key, value in cells if key == "Record") == Counter(expected_records)
    parsed_xlsx_manifest = [(parts[0], parts[1], parts[2], parts[3])
            for key, value in cells if key == "Source"
            for parts in [value.split("/")]]
    assert Counter(parsed_xlsx_manifest) == Counter(expected_manifest)
    parsed_xlsx_refs = [("manifest", *ref) for ref in parsed_xlsx_manifest]
    parsed_xlsx_refs += _parsed_detail_refs(cells)
    assert Counter(parsed_xlsx_refs) == Counter(expected_refs)
    assert meta == {**expected_metadata, "version": report.version}
    assert any(v == "=HYPERLINK(\"https://example.test\") " + "x" * 800
               for row in cells for v in row)
    assert not any(c.data_type == "f" or c.hyperlink for row in wb["Content"] for c in row)
    after = hashlib.sha256(canonical_bytes(report.model_dump(mode="json"))).hexdigest()
    assert before == after
    shape = next(s for outcome in MATRIX["outcomes"] if outcome["outcomeId"] == "EXPORTED"
                 for s in outcome["shapes"] if s["shapeId"] ==
                 f"P5A_EXPORTED_{status}_{'CURRENT' if current else 'HISTORICAL'}")
    for fmt, actual_records, actual_refs, actual_meta in (
        ("pdf", parsed_records, parsed_refs, pdf_meta),
        ("xlsx", [(value.split("/", 1)[0], value.split("/", 1)[1])
                  for key, value in cells if key == "Record"], parsed_xlsx_refs,
         {**meta, "version": str(meta["version"])}),
    ):
        receipt = {
            "outcome": "EXPORTED", "http_status": 200, "format": fmt,
            "view": "shift-summary", "report_status": status, "historical": not current,
            "attachment": True, "expected_record_count": len(expected_records),
            "rendered_record_count": len(actual_records),
            "expected_reference_count": len(expected_refs),
            "rendered_reference_count": len(actual_refs),
            "expected_record_membership_digest": _membership_digest(expected_records),
            "rendered_record_membership_digest": _membership_digest(actual_records),
            "expected_reference_membership_digest": _membership_digest(expected_refs),
            "rendered_reference_membership_digest": _membership_digest(actual_refs),
            "report_id": str(report.report_id), "shift_id": str(report.shift_id),
            "version": report.version, "is_current": current,
            "generated_from_cutoff": report.generated_from_cutoff.isoformat(),
            "snapshot_digest": report.content.snapshot_digest,
            "template_version": "p5a-1",
            "expected_metadata_digest": _metadata_digest(expected_metadata),
            "rendered_metadata_digest": _metadata_digest(actual_meta),
            "before_state_digest": before, "after_state_digest": after,
            "report_type": "END_SHIFT", "status_label": actual_meta["status"],
            "historical_label": actual_meta.get("Historical") == "Yes",
        }
        assert matches_shape_exactly(shape, receipt, MATRIX)
        if fmt == "pdf":
            mutations = generate_mutations(MATRIX, shape["shapeId"], receipt)
            assert mutations
            assert all(not matches_shape_exactly(shape, m.payload, MATRIX) for m in mutations)
    (PROOF_DIR / f"{name}.json").write_text(json.dumps({
        "report_id": str(report.report_id), "shift_id": str(report.shift_id),
        "snapshot_digest": report.content.snapshot_digest, "status": status,
        "is_current": current, "expected_records": sorted(expected_records),
        "expected_references": sorted(expected_refs),
        "before_state_digest": before, "after_state_digest": after,
        "pdf_sha256": hashlib.sha256(pdf).hexdigest(),
        "xlsx_sha256": hashlib.sha256(xlsx).hexdigest(),
    }, indent=2), encoding="utf-8")


@pytest.mark.parametrize("view,section", [
    ("customer-report", "customer_requests"),
    ("incident-report", "incidents"),
    ("handover-report", "handovers"),
])
def test_nonempty_focused_views_exclude_other_sections(view, section):
    report = full_report()
    selected = next(s for s in report.content.sections if s.section_type == section)
    expected_record = (selected.records[0]["record_type"], selected.records[0]["record_id"])
    expected_ref = next(m for m in report.content.source_manifest
                        if str(m.record_id) == expected_record[1])
    pdf = render_bounded(report, view, "pdf")
    text = "\n".join(page.get_text() for page in fitz.open(stream=pdf, filetype="pdf"))
    assert set(re.findall(r"Record: (\w+)/([0-9a-f-]{36})", text)) == {expected_record}
    assert expected_ref.source_digest in text.replace("\n", "")
    assert f"view: {view}" in text
    xlsx = render_bounded(report, view, "xlsx")
    wb = load_workbook(BytesIO(xlsx), read_only=True)
    rows = list(wb["Content"].values)
    assert {value for key, value in rows if key == "Record"} == {"/".join(expected_record)}
    assert {value for key, value in rows if key == "Source"} == {
        "/".join((expected_ref.record_type, str(expected_ref.record_id),
                  str(expected_ref.source_version), expected_ref.source_digest))
    }
    assert dict(list(wb["Metadata"].values)[1:])["view"] == view
