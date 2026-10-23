"""Real parsed output smoke and refusal proof for persisted empty snapshots."""

import hashlib
import json
from datetime import UTC, datetime
from io import BytesIO
from uuid import uuid4

import fitz
import pytest
from openpyxl import load_workbook
from operations_domain.models import Report, ReportContent, ReportSection, ReportSourceRef
from operations_domain.report_digest import aggregate_digest, canonical_bytes
from reporting_engine import ExportFailure, render_bounded
from workspace_api.application.report_snapshot import build_snapshot


def _report():
    shift_id = uuid4()
    content = build_snapshot(
        shift_id=shift_id, events=[], corrections=[], tasks=[],
        customer_requests=[], incidents=[], handovers=[],
    )
    return Report(shift_id=shift_id, content=content,
                  generated_from_cutoff=datetime.now(UTC))


@pytest.mark.parametrize("view", ["shift-summary", "customer-report", "incident-report", "handover-report"])
def test_actual_empty_artifact_parses_with_lineage(view):
    report = _report()
    pdf = render_bounded(report, view, "pdf")
    assert pdf.startswith(b"%PDF")
    document = fitz.open(stream=pdf, filetype="pdf")
    text = "\n".join(page.get_text() for page in document)
    assert str(report.report_id) in text and report.content.snapshot_digest in text
    assert "No records" in text and view in text
    xlsx = render_bounded(report, view, "xlsx")
    wb = load_workbook(BytesIO(xlsx), read_only=True, data_only=False)
    assert wb.sheetnames == ["Metadata", "Content", "Date values"]
    metadata = dict(wb["Metadata"].values)
    assert metadata["report_id"] == str(report.report_id)
    assert metadata["snapshot_digest"] == report.content.snapshot_digest
    assert ("Empty", "No records") in list(wb["Content"].values)


def test_wrong_but_valid_aggregate_digest_refused_before_bytes():
    report = _report()
    report.content.snapshot_digest = "a" * 64
    with pytest.raises(ExportFailure, match="EXPORT_INVALID_SNAPSHOT"):
        render_bounded(report, "shift-summary", "pdf")


def test_extracted_aggregate_helper_is_byte_identical_to_p2r_recipe():
    report = _report()
    content = report.content
    old_payload = {
        "schema_version": "1.0", "shift_id": str(report.shift_id),
        "report_type": "END_SHIFT",
        "sections": [s.model_dump(mode="json") for s in content.sections],
        "source_manifest": [m.model_dump(mode="json") for m in content.source_manifest],
    }
    old_bytes = json.dumps(old_payload, sort_keys=True, separators=(",", ":"),
                           ensure_ascii=False).encode("utf-8")
    assert content.snapshot_digest == hashlib.sha256(old_bytes).hexdigest()
    assert old_bytes == canonical_bytes(old_payload)


def test_reference_multiset_digest_rejects_substitution_and_duplicate_omission():
    def digest(items):
        return hashlib.sha256(canonical_bytes(sorted(items))).hexdigest()

    refs = [("manifest", "Task", "id1"), ("evidence", "E1"),
            ("evidence", "E1"), ("handover", "Task", "id1")]
    same_count_substitution = [*refs[:3], ("handover", "Task", "id2")]
    duplicate_omission = [refs[0], refs[1], refs[3], refs[3]]
    assert len(refs) == len(same_count_substitution) == len(duplicate_omission)
    assert digest(refs) != digest(same_count_substitution)
    assert digest(refs) != digest(duplicate_omission)


def test_actual_record_and_formula_text_are_literal_in_both_formats():
    import hashlib

    report = _report()
    record_id = uuid4()
    record = {
        "record_type": "Task", "record_id": str(record_id), "shift_id": str(report.shift_id),
        "title": "Báo cáo ế =1+1", "description": "=HYPERLINK(\"https://example.test\")",
        "status": "OPEN", "owner_id": None, "due_at": None, "risk_class": "R1",
        "state": "CONFIRMED", "version": 1, "created_at": "2026-01-01T00:00:00Z",
        "evidence": [],
    }
    ref = ReportSourceRef(record_type="Task", record_id=record_id, source_version=1,
                          source_digest=hashlib.sha256(canonical_bytes(record)).hexdigest())
    sections = [ReportSection(section_type=s.section_type,
                              records=[record] if s.section_type == "tasks" else [])
                for s in report.content.sections]
    content = ReportContent(sections=sections, source_manifest=[ref], snapshot_digest="0" * 64)
    content.snapshot_digest = aggregate_digest(report.shift_id, report.report_type, content)
    old_payload = {
        "schema_version": "1.0", "shift_id": str(report.shift_id),
        "report_type": "END_SHIFT",
        "sections": [s.model_dump(mode="json") for s in content.sections],
        "source_manifest": [m.model_dump(mode="json") for m in content.source_manifest],
    }
    old_bytes = json.dumps(old_payload, sort_keys=True, separators=(",", ":"),
                           ensure_ascii=False).encode("utf-8")
    assert content.snapshot_digest == hashlib.sha256(old_bytes).hexdigest()
    assert old_bytes == canonical_bytes(old_payload)
    report.content = content
    pdf = render_bounded(report, "shift-summary", "pdf")
    text = "\n".join(p.get_text() for p in fitz.open(stream=pdf, filetype="pdf"))
    assert str(record_id) in text and str(ref.source_digest) in text.replace("\n", "")
    assert "Báo cáo" in text
    xlsx = render_bounded(report, "shift-summary", "xlsx")
    wb = load_workbook(BytesIO(xlsx), data_only=False)
    cells = [c for row in wb["Content"] for c in row]
    assert any(c.value == record["description"] and c.data_type == "s" for c in cells)
    assert any(str(record_id) in str(c.value) for c in cells)
