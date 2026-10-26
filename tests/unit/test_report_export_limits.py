"""Bounded renderer refusals, including a child stalled after partial output."""

import multiprocessing as mp
import os
import time

import pytest
from reporting_engine.limits import validate_text
from reporting_engine.model import ExportFailure
from reporting_engine.pdf import render_pdf
from reporting_engine.worker import _run_isolated
from reporting_engine.xlsx import render_xlsx


def _partial_then_stall(shared, size, status, marker):
    shared[0] = ord("X")
    size.value = 1
    status.value = 1
    marker.write_text(str(os.getpid()), encoding="utf-8")
    time.sleep(20)


def test_partial_worker_output_is_killed_and_never_returned(tmp_path):
    marker = tmp_path / "after-partial-write.pid"
    start = time.monotonic()
    with pytest.raises(ExportFailure, match="EXPORT_LIMIT_EXCEEDED"):
        _run_isolated(_partial_then_stall, (marker,), 3)
    assert marker.is_file(), "child must reach the partial write before deadline"
    pid = int(marker.read_text(encoding="utf-8"))
    assert pid not in {child.pid for child in mp.active_children()}
    assert 3 <= time.monotonic() - start < 6


def test_xlsx_utf16_cell_limit_and_pdf_independence():
    from io import BytesIO

    from openpyxl import load_workbook

    near = "a" * 32_765 + "\U0001f600"
    validate_text(near, pdf=False)
    data = render_xlsx([], [("value", near)])
    wb = load_workbook(BytesIO(data), read_only=True, data_only=False)
    assert list(wb["Content"].values)[1][1] == near
    with pytest.raises(ExportFailure, match="EXPORT_LIMIT_EXCEEDED"):
        render_xlsx([], [("value", near + "b")])
    validate_text("a" * 32_768, pdf=True)


@pytest.mark.parametrize("character", ["\x00", "\ud800", "\ufffe"])
def test_invalid_xml_text_fails_controlled(character):
    with pytest.raises(ExportFailure, match="EXPORT_UNSUPPORTED_TEXT"):
        render_xlsx([], [("value", character)])


def test_formula_payload_stays_literal_and_rows_are_bounded():
    with pytest.raises(ExportFailure, match="EXPORT_LIMIT_EXCEEDED"):
        render_xlsx([], [("record", "x")] * 19_999)
    data = render_xlsx([], [("record", "=1+1")])
    assert data.startswith(b"PK")


def test_xlsx_aggregate_row_cap_at_boundary(monkeypatch):
    from io import BytesIO

    from openpyxl import load_workbook
    from reporting_engine import xlsx

    monkeypatch.setattr(xlsx, "MAX_ROWS", 6)
    data = render_xlsx([("view", "summary")], [("record", "one"), ("record", "two")])
    wb = load_workbook(BytesIO(data), read_only=True)
    assert sum(ws.max_row for ws in wb) == 6
    with pytest.raises(ExportFailure, match="EXPORT_LIMIT_EXCEEDED"):
        render_xlsx([("view", "summary")], [("record", "one")] * 3)


def test_actual_20000_row_xlsx_serializes_without_truncation():
    from io import BytesIO

    from openpyxl import load_workbook

    data = render_xlsx([], [("record", "literal") for _ in range(19_997)])
    wb = load_workbook(BytesIO(data), read_only=True)
    assert sum(ws.max_row for ws in wb) == 20_000
    with pytest.raises(ExportFailure, match="EXPORT_LIMIT_EXCEEDED"):
        render_xlsx([], [("record", "literal") for _ in range(19_998)])


def test_xlsx_literal_payloads_roundtrip_including_carriage_return():
    from io import BytesIO

    from openpyxl import load_workbook

    values = ["=1+1", "+2", "-3", "@name", "\tleading", "\rvalue",
              "https://example.test/path"]
    data = render_xlsx([], [("literal", value) for value in values])
    wb = load_workbook(BytesIO(data), data_only=False)
    cells = [row[1] for row in list(wb["Content"])[1:]]
    assert [cell.value for cell in cells] == values
    assert all(cell.data_type == "s" and cell.hyperlink is None for cell in cells)
    assert not wb._external_links


def test_xlsx_controlled_version_and_utc_derivative_types():
    from datetime import datetime
    from io import BytesIO

    from openpyxl import load_workbook

    iso = "2026-01-01T07:00:00.123456+07:00"
    data = render_xlsx([("version", "3"), ("generated_from_cutoff", iso)], [
        ("Record", "Task/abc"), ("previous_version", "1"),
        ("new_version", "2"), ("created_at", iso),
    ])
    wb = load_workbook(BytesIO(data), data_only=False)
    meta = {row[0].value: row[1] for row in list(wb["Metadata"])[1:]}
    content = {row[0].value: row[1] for row in list(wb["Content"])[1:]}
    assert meta["version"].value == 3 and meta["version"].data_type == "n"
    assert content["previous_version"].value == 1
    assert content["new_version"].value == 2
    assert meta["generated_from_cutoff"].value == iso
    assert content["created_at"].value == iso
    dates = list(wb["Date values"].values)
    assert dates[0] == ("Context", "Field", "Canonical ISO",
                        "Excel datetime derivative (millisecond precision)",
                        "Timezone / representation note")
    assert dates[1][:4] == ("report", "generated_from_cutoff", iso,
                            datetime(2026, 1, 1, 0, 0, 0, 123000))
    assert dates[2][:4] == ("Task/abc", "created_at", iso,
                            datetime(2026, 1, 1, 0, 0, 0, 123000))
    assert all("converted to UTC" in row[4] for row in dates[1:])
    assert wb["Date values"]["D2"].data_type == "d"


def test_xlsx_naive_timestamp_keeps_wall_clock_without_utc_claim():
    from datetime import datetime
    from io import BytesIO

    from openpyxl import load_workbook

    iso = "2026-01-01T07:08:09.123456"
    wb = load_workbook(BytesIO(render_xlsx([("generated_from_cutoff", iso)], [])))
    assert dict(list(wb["Metadata"].values)[1:])["generated_from_cutoff"] == iso
    context, field, original, derivative, note = list(wb["Date values"].values)[1]
    assert (context, field, original) == ("report", "generated_from_cutoff", iso)
    assert derivative == datetime(2026, 1, 1, 7, 8, 9, 123000)
    assert "Timezone unspecified" in note and "no conversion" in note


@pytest.mark.parametrize("value", ["999999999999999", "-999999999999999", "0", "-1"])
def test_xlsx_safe_integer_stays_numeric_exact(value):
    from io import BytesIO

    from openpyxl import load_workbook

    wb = load_workbook(BytesIO(render_xlsx([("version", value)], [])))
    cell = wb["Metadata"]["B2"]
    assert cell.data_type == "n" and cell.value == int(value)
    assert cell.comment is None


@pytest.mark.parametrize("value", ["1000000000000000", "-1000000000000000",
                                      "18014398509481985", "-18014398509481985"])
def test_xlsx_large_integer_is_annotated_exact_text(value):
    from io import BytesIO

    from openpyxl import load_workbook

    wb = load_workbook(BytesIO(render_xlsx([("version", value)], [])))
    cell = wb["Metadata"]["B2"]
    assert cell.data_type == "s" and cell.value == value
    assert cell.number_format == "@"
    assert cell.comment and "Exact integer text" in cell.comment.text


@pytest.mark.parametrize("iso,fragment", [
    ("1899-12-31T23:59:59Z", "outside Excel date range"),
    ("0001-01-01T00:00:00+07:00", "outside Excel date range"),
    ("9999-12-31T23:59:59-07:00", "outside Excel date range"),
])
def test_xlsx_unrepresentable_derivative_retains_canonical_iso(iso, fragment):
    from io import BytesIO

    from openpyxl import load_workbook

    wb = load_workbook(BytesIO(render_xlsx([("generated_from_cutoff", iso)], [])))
    assert dict(list(wb["Metadata"].values)[1:])["generated_from_cutoff"] == iso
    context, field, original, derivative, note = list(wb["Date values"].values)[1]
    assert (context, field, original) == ("report", "generated_from_cutoff", iso)
    assert derivative is None and fragment in note


def test_xlsx_date_derivative_rows_count_toward_cap(monkeypatch):
    from reporting_engine import xlsx

    monkeypatch.setattr(xlsx, "MAX_ROWS", 5)
    one_date = [("generated_from_cutoff", "2026-01-01T00:00:00Z")]
    render_xlsx(one_date, [])  # Three sheet headers, primary row, derivative row.
    with pytest.raises(ExportFailure, match="EXPORT_LIMIT_EXCEEDED"):
        render_xlsx(one_date, [("created_at", "2026-01-01T00:00:00Z")])


def test_pdf_missing_font_and_unknown_glyph(monkeypatch, tmp_path):
    from reporting_engine import limits

    with pytest.raises(ExportFailure, match="EXPORT_UNSUPPORTED_TEXT"):
        render_pdf([("view", "shift-summary")], [("value", "\U0001f680")])
    monkeypatch.setattr(limits, "FONT_PATH", tmp_path / "missing.ttf")
    with pytest.raises(ExportFailure, match="EXPORT_UNSUPPORTED_TEXT"):
        render_pdf([("view", "shift-summary")], [("value", "hello")])


def test_pdf_page_and_serialized_byte_caps(monkeypatch):
    import fitz
    from reporting_engine import pdf

    metadata = [("view", "shift-summary")]
    rows = [("value", "x" * 12_000)]
    monkeypatch.setattr(pdf, "MAX_PAGES", 1)
    one_page = render_pdf(metadata, [("value", "x")])
    assert len(fitz.open(stream=one_page, filetype="pdf")) == 1
    with pytest.raises(ExportFailure, match="EXPORT_LIMIT_EXCEEDED"):
        render_pdf(metadata, rows)
    monkeypatch.setattr(pdf, "MAX_PAGES", 512)
    monkeypatch.setattr(pdf, "MAX_BYTES", 1)
    with pytest.raises(ExportFailure, match="EXPORT_LIMIT_EXCEEDED"):
        render_pdf(metadata, [("value", "x")])
