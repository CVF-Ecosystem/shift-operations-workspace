"""Typed XLSX presentation with literal user text and bounded output."""

import io
from datetime import UTC, datetime

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font

from .limits import MAX_BYTES, validate_text
from .model import ExportFailure
from .xlsx_xml import preserve_carriage_returns

MAX_ROWS = 20_000
NUMERIC_FIELDS = {"version", "previous_version", "new_version"}
DATE_FIELDS = {
    "generated_from_cutoff", "starts_at", "ends_at", "created_at", "due_at",
    "received_at", "promised_at", "reviewed_at", "acknowledged_at",
}
EXCEL_MIN_DATE = datetime(1900, 1, 1)


def _integer_cell(value: str) -> tuple[int | str, bool]:
    try:
        parsed = int(value)
    except ValueError as exc:
        raise ExportFailure("EXPORT_INVALID_SNAPSHOT") from exc
    if len(str(abs(parsed))) <= 15:
        return parsed, False
    return value, True


def _date_rows(metadata: list[tuple[str, str]], rows: list[tuple[str, str]]) -> list[tuple]:
    result = []
    context = "report"
    for field, value in metadata + rows:
        if field == "Section":
            context = f"section:{value}"
        elif field == "Record":
            context = value
        if field not in DATE_FIELDS or value in {"", "null"}:
            continue
        try:
            moment = datetime.fromisoformat(value)
        except ValueError as exc:
            raise ExportFailure("EXPORT_INVALID_SNAPSHOT") from exc
        if moment.tzinfo is None or moment.utcoffset() is None:
            derivative = moment
            note = "Timezone unspecified; wall-clock only; no conversion"
        else:
            note = "Explicit timezone; converted to UTC"
            try:
                derivative = moment.astimezone(UTC).replace(tzinfo=None)
            except OverflowError:
                derivative = None
        if derivative is None or derivative < EXCEL_MIN_DATE:
            derivative = None
            note += "; derivative outside Excel date range"
        else:
            derivative = derivative.replace(
                microsecond=(derivative.microsecond // 1000) * 1000,
            )
            note += "; millisecond precision (truncated)"
        result.append((context, field, value, derivative, note))
    return result


def render_xlsx(metadata: list[tuple[str, str]], rows: list[tuple[str, str]]) -> bytes:
    dates = _date_rows(metadata, rows)
    if len(metadata) + len(rows) + len(dates) + 3 > MAX_ROWS:
        raise ExportFailure("EXPORT_LIMIT_EXCEEDED")
    wb = Workbook()
    meta = wb.active
    meta.title = "Metadata"
    content = wb.create_sheet("Content")
    date_sheet = wb.create_sheet("Date values")
    wrapped = Alignment(wrap_text=True, vertical="top")
    for sheet, pairs in ((meta, metadata), (content, rows)):
        sheet.append(("Field", "Value"))
        for row_number, (field, value) in enumerate(pairs, start=2):
            validate_text(field, pdf=False)
            validate_text(value, pdf=False)
            if field in NUMERIC_FIELDS:
                cell_value, exact_text = _integer_cell(value)
            else:
                cell_value, exact_text = value, False
            sheet.append((field, cell_value))
            for column in (1, 2):
                cell = sheet.cell(row_number, column)
                if column == 1 or field not in NUMERIC_FIELDS or exact_text:
                    cell.data_type = "s"
                cell.alignment = wrapped
            if exact_text:
                value_cell = sheet.cell(row_number, 2)
                value_cell.number_format = "@"
                value_cell.comment = Comment(
                    "Exact integer text; exceeds Excel's reliable 15-digit numeric precision.",
                    "reporting-engine",
                )
        for cell in (sheet.cell(1, 1), sheet.cell(1, 2)):
            cell.font = Font(bold=True)
        sheet.column_dimensions["A"].width = 28
        sheet.column_dimensions["B"].width = 90

    date_sheet.append(("Context", "Field", "Canonical ISO",
                       "Excel datetime derivative (millisecond precision)", "Timezone / representation note"))
    for row_number, values in enumerate(dates, start=2):
        for value in values[:3]:
            validate_text(value, pdf=False)
        date_sheet.append(values)
        for column in (1, 2, 3):
            date_sheet.cell(row_number, column).data_type = "s"
        if values[3] is not None:
            date_sheet.cell(row_number, 4).number_format = "yyyy-mm-dd hh:mm:ss.000"
    for cell in date_sheet[1]:
        cell.font = Font(bold=True)
    for column, width in (("A", 55), ("B", 28), ("C", 36), ("D", 42), ("E", 70)):
        date_sheet.column_dimensions[column].width = width

    output = io.BytesIO()
    wb.save(output)
    data = output.getvalue()
    if any("\r" in value for _, value in metadata + rows):
        data = preserve_carriage_returns(data)
    if len(data) > MAX_BYTES:
        raise ExportFailure("EXPORT_LIMIT_EXCEEDED")
    return data
